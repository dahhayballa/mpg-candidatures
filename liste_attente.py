#!/usr/bin/env python3
"""
MPG — Liste d'attente par spécialité
10% du quota par filière
"""
import sqlite3, sys, os
sys.path.append(os.path.dirname(__file__))
import config

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    os.system("pip install openpyxl")
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB_PATH = config.DB_PATH

QUOTAS = {
    "HSE":                                40,
    "Électricité industrielle":           30,
    "Tuyauterie industrielle":            30,
    "Construction métallique et soudure": 40,
    "Maintenance industrielle":           20,
    "Techniques minières":                20,
    "Opérations pétrolières et gazières": 20,
}

conn = sqlite3.connect(DB_PATH)

# Styles
HDR_FILL = PatternFill("solid", fgColor="1E3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11, name="Arial")
GREEN    = PatternFill("solid", fgColor="D5F5E3")
ORANGE   = PatternFill("solid", fgColor="FDEBD0")
YELLOW   = PatternFill("solid", fgColor="FEF3C7")
WHITE    = PatternFill("solid", fgColor="FFFFFF")
GRAY     = PatternFill("solid", fgColor="F4F6F7")
thin     = Side(style='thin', color='BDC3C7')
border   = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
wb.remove(wb.active)

print("=" * 65)
print("  EETFP-MPG — Listes d'attente (10% par filière)")
print("=" * 65)

summary_data = []

for specialty, quota in QUOTAS.items():

    # 10% du quota — minimum 1
    nb_attente = max(1, round(quota * 0.10))

    # Retenus pour calculer le cutoff
    retenus = conn.execute("""
        SELECT COUNT(*), MIN(score_total)
        FROM candidates
        WHERE specialty=?
        AND hors_delai=0 AND retenu=1
    """, (specialty,)).fetchone()
    nb_retenus  = retenus[0] or 0
    score_cutoff = retenus[1] or 0

    # Liste d'attente — meilleurs non retenus
    attente = conn.execute("""
        SELECT
            id, name, specialty, score_total,
            score_niveau, score_experience, score_motivation,
            score_adequation, score_dossier,
            mention, source, first_date,
            justif_niveau, note_globale
        FROM candidates
        WHERE specialty=?
        AND hors_delai=0
        AND (retenu=0 OR retenu IS NULL)
        AND score_total >= 50
        ORDER BY score_total DESC, first_date ASC
        LIMIT ?
    """, (specialty, nb_attente)).fetchall()

    libres = quota - nb_retenus

    print(f"\n📋 {specialty}")
    print(f"   Quota: {quota} | Retenus: {nb_retenus} | "
          f"Places libres: {libres} | Liste attente: {nb_attente}")

    # Feuille Excel
    short = specialty[:28].replace('/', '-')
    ws = wb.create_sheet(title=short)

    # Titre
    ws.merge_cells('A1:L1')
    ws['A1'] = f"LISTE D'ATTENTE — {specialty.upper()}"
    ws['A1'].font = Font(bold=True, color="FFFFFF", size=13, name="Arial")
    ws['A1'].fill = HDR_FILL
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    # Infos
    ws.merge_cells('A2:L2')
    ws['A2'] = (f"Quota: {quota} places | Retenus: {nb_retenus} | "
                f"Score cutoff: {score_cutoff}/100 | "
                f"Liste d'attente: {nb_attente} candidats (10%)")
    ws['A2'].font = Font(italic=True, color="7F8C8D", size=10, name="Arial")
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 18

    # En-têtes
    headers = ["Rang","Nom complet","Total/100","N/25","E/20",
               "M/20","A/20","D/10","Mention","Source","Date","Note IA"]
    widths  = [6, 32, 10, 7, 7, 7, 7, 7, 12, 10, 14, 42]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = HDR_FONT
        c.fill = PatternFill("solid", fgColor="2D6A9F")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[3].height = 24

    # Lignes de données
    for rang, row in enumerate(attente, 1):
        (cid, name, spec, total, n, e, m, a, d,
         mention, source, fdate, justif, note) = row

        r = rang + 3
        fill = GRAY if rang % 2 == 0 else WHITE

        vals = [
            rang,
            name or "",
            total or 0,
            n or 0, e or 0, m or 0, a or 0, d or 0,
            mention or "",
            "Papier" if source == "papier" else "Email",
            str(fdate or "")[:10],
            note or ""
        ]

        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = fill
            c.border = border
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(
                vertical='center',
                horizontal='center' if col in [1,3,4,5,6,7,8,10] else 'left',
                wrap_text=(col == 12)
            )

        # Couleur mention
        mention_fill = {
            "Excellent": GREEN,
            "Bon":       PatternFill("solid", fgColor="DBEAFE"),
            "Moyen":     YELLOW,
        }.get(mention, WHITE)
        ws.cell(r, 9).fill = mention_fill

        # Couleur score
        if total and total >= 65:
            ws.cell(r, 3).fill = GREEN
        elif total and total >= 50:
            ws.cell(r, 3).fill = YELLOW

        ws.row_dimensions[r].height = 40

        print(f"   #{rang} {str(name or '')[:38]:<38} {total:>5}/100 — {mention}")

    ws.freeze_panes = 'A4'

    summary_data.append({
        "specialty": specialty, "quota": quota,
        "retenus": nb_retenus, "libres": libres,
        "nb_attente": nb_attente, "attente": len(attente)
    })

# ── Feuille Résumé ──────────────────────────────────────────
ws_r = wb.create_sheet(title="RÉSUMÉ", index=0)

ws_r.merge_cells('A1:G1')
ws_r['A1'] = "EETFP-MPG — RÉSUMÉ LISTES D'ATTENTE — FQ2025-2026"
ws_r['A1'].font = Font(bold=True, color="FFFFFF", size=13, name="Arial")
ws_r['A1'].fill = HDR_FILL
ws_r['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws_r.row_dimensions[1].height = 32

for col, h in enumerate(
    ["Filière","Quota","Retenus","Places libres","Taille liste attente","Candidats trouvés","Statut"], 1
):
    c = ws_r.cell(2, col, h)
    c.font = HDR_FONT
    c.fill = PatternFill("solid", fgColor="2D6A9F")
    c.alignment = Alignment(horizontal='center', wrap_text=True)
    c.border = border
ws_r.row_dimensions[2].height = 28

ws_r.column_dimensions['A'].width = 35
for col in ['B','C','D','E','F','G']:
    ws_r.column_dimensions[col].width = 16

conn = sqlite3.connect(DB_PATH)
for i, s in enumerate(summary_data, 3):
    # Recalcul direct depuis DB avec specialty exacte
    rn = conn.execute("SELECT COUNT(*) FROM candidates WHERE specialty=? AND hors_delai=0 AND retenu=1",
                      (s["specialty"],)).fetchone()[0]
    an = conn.execute("SELECT COUNT(*) FROM candidates WHERE specialty=? AND hors_delai=0 AND (retenu=0 OR retenu IS NULL) AND score_total>=50",
                      (s["specialty"],)).fetchone()[0]
    libres = s["quota"] - rn
    fill = ORANGE if libres > 0 else GREEN
    statut = f"⚠️ {libres} places libres" if libres > 0 else "✅ Complet"
    vals = [s["specialty"], s["quota"], rn, libres, s["nb_attente"], an, statut]
    for col, val in enumerate(vals, 1):
        c = ws_r.cell(i, col, val)
        c.fill = fill
        c.border = border
        c.font = Font(name="Arial", size=11)
        c.alignment = Alignment(
            horizontal='left' if col == 1 else 'center',
            vertical='center'
        )
    ws_r.row_dimensions[i].height = 22

# Total
tot_row = len(summary_data) + 3
for col, val in enumerate([
    "TOTAL",
    sum(s["quota"] for s in summary_data),
    sum(s["retenus"] for s in summary_data),
    sum(s["libres"] for s in summary_data),
    sum(s["nb_attente"] for s in summary_data),
    sum(s["attente"] for s in summary_data),
    ""
], 1):
    c = ws_r.cell(tot_row, col, val)
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    c.fill = HDR_FILL
    c.border = border
    c.alignment = Alignment(horizontal='left' if col==1 else 'center', vertical='center')
ws_r.row_dimensions[tot_row].height = 25
ws_r.freeze_panes = 'A3'

conn.close()

from datetime import datetime
fname = f"Liste_Attente_MPG_FQ2026_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
wb.save(fname)

print(f"\n{'='*65}")
print(f"  ✅ Fichier : {fname}")
print(f"  Total liste attente : {sum(s['nb_attente'] for s in summary_data)} candidats")
print(f"  (10% de chaque quota)")
print(f"{'='*65}")