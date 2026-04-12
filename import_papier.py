#!/usr/bin/env python3
"""
════════════════════════════════════════════════════════════════
  MPG — Import Candidats Papier (v2 — colonnes corrigées)
  Usage : python import_papier.py
════════════════════════════════════════════════════════════════
"""
import sqlite3, sys, os
sys.path.append(os.path.dirname(__file__))
import config

try:
    import openpyxl
except ImportError:
    os.system("pip install openpyxl")
    import openpyxl

DB_PATH    = config.DB_PATH
EXCEL_FILE = "Liste_FQ2026__1_.xlsx"

SHEET_TO_SPECIALTY = {
    "Opération pétrolières et gazièr": "Opérations pétrolières et gazières",
    "Electricité Industrielle ":        "Électricité industrielle",
    "HSE":                              "HSE",
    "Maintenance Industrielle ":        "Maintenance industrielle",
    "Construction métallique et Soud":  "Construction métallique et soudure",
    "Technique minière ":               "Techniques minières",
    "Tuyautérie":                       "Tuyauterie industrielle",
}

def get_score_niveau(diplome_str):
    """Retourne (score, label) selon grille officielle."""
    if not diplome_str:
        return 0, "Inconnu"
    d = str(diplome_str).lower().strip()
    if any(k in d for k in ["master","ingénieur","ingenieur","doctorat"]):
        return 25, "Master/Ingénieur"
    if any(k in d for k in ["licence","bac+3","l3"]):
        return 20, "Licence"
    if any(k in d for k in ["bts","bac+2","b.t.s"]):
        return 20, "BTS"
    if any(k in d for k in ["bt ","b.t ","brevet tech"]) or d == "bt":
        return 15, "BT"
    if "cap" in d or "certificat d'aptitude" in d:
        return 10, "CAP"
    if d.startswith("cc") or "cours comp" in d:
        return 5, "CC"
    if "bac" in d or "baccal" in d:
        return 0, "Bac"
    # Formations spécialisées non reconnues → CAP
    if any(k in d for k in ["soudure","soudeur","formation","sc:","om"]):
        return 10, "Formation/CAP"
    if "????" in d or d == "":
        return 0, "Inconnu"
    return 0, "Inconnu"


def add_columns(conn):
    cols = [r[1] for r in conn.execute("PRAGMA table_info(candidates)").fetchall()]
    for col, typ in [("source","TEXT"), ("nni","TEXT"), ("telephone","TEXT")]:
        if col not in cols:
            conn.execute(f"ALTER TABLE candidates ADD COLUMN {col} {typ}")
    conn.commit()


def import_sheet(ws, specialty, conn):
    rows = list(ws.iter_rows(values_only=True))

    # Trouver la ligne de données (après "N°")
    data_start = None
    for i, row in enumerate(rows):
        vals = [str(v or "").strip() for v in row]
        if any("N°" in v or "Noms" in v for v in vals):
            data_start = i + 1
            break

    if data_start is None:
        print(f"   ⚠ En-tête non trouvé")
        return []

    # Structure colonnes (basée sur l'analyse réelle) :
    # [0]=None [1]=N° [2]=Nom [3]=NNI [4]=Tel [5]=None
    # [6]=Diplôme [7]=Niveau_score [8]=Exp [9]=Motiv
    # [10]=Adeq [11]=Dossier [12]=Dispo [13]=Total

    imported = []
    for row in rows[data_start:]:
        if not any(v is not None for v in row):
            continue
        if row[1] is None:
            continue

        try:
            nom     = str(row[2] or "").strip()
            nni     = str(row[3] or "").strip()
            tel     = str(row[4] or "").strip()
            diplome = str(row[6] or "").strip()

            if not nom or nom == "None" or len(nom) < 3:
                continue

            # Scores de la commission (E, M, A, D)
            score_exp  = float(row[8]  or 0)
            score_mot  = float(row[9]  or 0)
            score_adeq = float(row[10] or 0)
            score_dos  = float(row[11] or 0)

            # Bornage
            score_exp  = max(0, min(20, score_exp))
            score_mot  = max(0, min(20, score_mot))
            score_adeq = max(0, min(20, score_adeq))
            score_dos  = max(0, min(10, score_dos))

            # Recalcul niveau avec grille officielle
            score_niveau, niveau_label = get_score_niveau(diplome)
            score_dispo = 5.0

            total = round(score_niveau + score_exp + score_mot +
                         score_adeq + score_dos + score_dispo, 1)

            mention = ("Excellent" if total >= 80 else
                      "Bon"        if total >= 65 else
                      "Moyen"      if total >= 50 else
                      "Non retenu")

            # Email fictif unique
            safe_nom = nom.lower().replace(' ','_').replace("'","")[:30]
            email = f"papier_{nni or safe_nom}@papier.eetfp-mpg.mr"

            # Vérifier doublon
            if conn.execute("SELECT id FROM candidates WHERE email_addr=?",
                           (email,)).fetchone():
                continue

            conn.execute("""
                INSERT INTO candidates (
                    email_addr, name, specialty, source, nni, telephone,
                    score_total, score_niveau, score_experience,
                    score_motivation, score_adequation, score_dossier,
                    score_disponibilite, mention, ia_scored,
                    has_cv, has_motivation, has_id, has_diplomas,
                    status, hors_delai, num_emails,
                    justif_niveau, note_globale,
                    first_date, last_date
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, (
                email, nom, specialty, "papier", nni, tel,
                total, score_niveau, score_exp,
                score_mot, score_adeq, score_dos,
                score_dispo, mention, 1,
                1, 1, 1, 1,
                "Complet", 0, 1,
                f"{diplome} → {niveau_label} = {score_niveau}/25 (grille officielle)",
                f"Dossier papier. Diplôme: {diplome}. Commission: E={score_exp} M={score_mot} A={score_adeq} D={score_dos}",
                "2026-03-31", "2026-03-31"
            ))

            imported.append({
                "nom": nom, "diplome": diplome,
                "niveau": niveau_label,
                "total": total, "mention": mention,
                "ancien_total": float(row[13] or 0)
            })

        except Exception as e:
            print(f"   ⚠ Erreur: {e} — ligne: {row[:7]}")
            continue

    conn.commit()
    return imported


def main():
    print("=" * 65)
    print("  MPG — Import Candidats Papier v2")
    print("=" * 65)

    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Fichier non trouvé : {EXCEL_FILE}")
        return

    conn = sqlite3.connect(DB_PATH)
    add_columns(conn)

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    total_imported = 0

    for sheet_name in wb.sheetnames:
        specialty = SHEET_TO_SPECIALTY.get(sheet_name, sheet_name.strip())
        ws = wb[sheet_name]

        print(f"\n📋 {specialty}")
        imported = import_sheet(ws, specialty, conn)

        for c in imported:
            diff = c['total'] - c['ancien_total']
            sign = f"+{diff:.1f}" if diff > 0 else f"{diff:.1f}"
            flag = "⬆" if diff > 0 else ("⬇" if diff < 0 else "=")
            print(f"   {flag} {c['nom'][:35]:<35} | {c['diplome'][:12]:<12} "
                  f"→ {c['niveau']:<18} | {c['total']:>5}/100 {c['mention']}"
                  f"  (Δ {sign} vs LJ)")
        if not imported:
            print(f"   → Aucun (vide ou déjà importé)")
        else:
            print(f"   ✅ {len(imported)} candidats")
        total_imported += len(imported)

    n_papier = conn.execute(
        "SELECT COUNT(*) FROM candidates WHERE source='papier'"
    ).fetchone()[0]

    print("\n" + "=" * 65)
    print(f"  ✅ {total_imported} candidats papier importés")
    print(f"  📊 Total papier en DB : {n_papier}")
    print()
    print("  Résumé par mention :")
    for r in conn.execute("""
        SELECT mention, COUNT(*) FROM candidates WHERE source='papier'
        GROUP BY mention ORDER BY
        CASE mention WHEN 'Excellent' THEN 1 WHEN 'Bon' THEN 2
        WHEN 'Moyen' THEN 3 ELSE 4 END
    """).fetchall():
        print(f"    {r[0]:<15} : {r[1]}")

    conn.close()
    print()
    print("  → Lancez : python generer_liste_finale.py")
    print("=" * 65)


if __name__ == "__main__":
    main()